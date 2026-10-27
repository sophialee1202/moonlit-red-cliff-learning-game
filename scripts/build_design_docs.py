from __future__ import annotations

import html
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "game-design-document.md"
HTML_OUTPUT = ROOT / "docs" / "game-design.html"
DOCX_OUTPUT = ROOT / "docs" / "game-design-plan-updated.docx"
FONT = "Microsoft JhengHei"


def cells(line: str) -> list[str]:
    return [part.strip() for part in line.strip().strip("|").split("|")]


def parse_markdown(lines: list[str]) -> list[tuple[str, object]]:
    events: list[tuple[str, object]] = []
    paragraph: list[str] = []
    index = 0

    def flush() -> None:
        nonlocal paragraph
        if paragraph:
            events.append(("paragraph", " ".join(part.strip() for part in paragraph)))
            paragraph = []

    while index < len(lines):
        line = lines[index]
        if not line.strip():
            flush()
            index += 1
            continue
        if line.startswith("# "):
            flush()
            events.append(("title", line[2:].strip()))
            index += 1
            continue
        if line.startswith("## "):
            flush()
            events.append(("heading1", line[3:].strip()))
            index += 1
            continue
        if line.startswith("### "):
            flush()
            events.append(("heading2", line[4:].strip()))
            index += 1
            continue
        if line.startswith(">"):
            flush()
            notes = []
            while index < len(lines) and lines[index].startswith(">"):
                notes.append(lines[index][1:].strip())
                index += 1
            events.append(("note", " ".join(notes)))
            continue
        if line.lstrip().startswith("|") and index + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{2,}", lines[index + 1]):
            flush()
            rows = [cells(line)]
            index += 2
            while index < len(lines) and lines[index].lstrip().startswith("|"):
                rows.append(cells(lines[index]))
                index += 1
            events.append(("table", rows))
            continue
        if re.match(r"^\s*-\s+", line):
            flush()
            items = []
            while index < len(lines) and re.match(r"^\s*-\s+", lines[index]):
                items.append(re.sub(r"^\s*-\s+", "", lines[index]).strip())
                index += 1
            events.append(("bullet", items))
            continue
        if re.match(r"^\s*\d+\.\s+", line):
            flush()
            items = []
            while index < len(lines) and re.match(r"^\s*\d+\.\s+", lines[index]):
                items.append(re.sub(r"^\s*\d+\.\s+", "", lines[index]).strip())
                index += 1
            events.append(("number", items))
            continue
        paragraph.append(line)
        index += 1
    flush()
    return events


INLINE = re.compile(r"(\*\*.+?\*\*|\*.+?\*|\[[^\]]+\]\(https?://[^)]+\))")


def inline_parts(value: str):
    cursor = 0
    for match in INLINE.finditer(value):
        if match.start() > cursor:
            yield "text", value[cursor : match.start()]
        token = match.group(0)
        if token.startswith("**"):
            yield "bold", token[2:-2]
        elif token.startswith("*"):
            yield "italic", token[1:-1]
        else:
            link = re.match(r"\[([^\]]+)\]\((https?://[^)]+)\)", token)
            yield "link", (link.group(1), link.group(2))
        cursor = match.end()
    if cursor < len(value):
        yield "text", value[cursor:]


def plain(value: str) -> str:
    output: list[str] = []
    for kind, item in inline_parts(value):
        if kind == "link":
            label, url = item
            output.append(f"{label}（{url}）")
        else:
            output.append(item)
    return "".join(output)


def html_inline(value: str) -> str:
    output: list[str] = []
    for kind, item in inline_parts(value):
        if kind == "link":
            label, url = item
            output.append(f'<a href="{html.escape(url, quote=True)}" target="_blank" rel="noopener noreferrer">{html.escape(label)}</a>')
        elif kind == "bold":
            output.append(f"<strong>{html.escape(item)}</strong>")
        elif kind == "italic":
            output.append(f"<em>{html.escape(item)}</em>")
        else:
            output.append(html.escape(item))
    return "".join(output)


def build_html(events: list[tuple[str, object]]) -> None:
    toc: list[tuple[str, str]] = []
    body: list[str] = []
    title = "月下赤壁 意象尋蹤 遊戲企劃書"
    section = 0
    for kind, value in events:
        if kind == "title":
            continue
        if kind == "heading1":
            section += 1
            anchor = f"section-{section}"
            toc.append((anchor, plain(value)))
            body.append(f'<h2 id="{anchor}">{html_inline(value)}</h2>')
        elif kind == "heading2":
            body.append(f"<h3>{html_inline(value)}</h3>")
        elif kind == "paragraph":
            body.append(f"<p>{html_inline(value)}</p>")
        elif kind == "note":
            body.append(f'<aside class="note">{html_inline(value)}</aside>')
        elif kind in ("bullet", "number"):
            tag = "ul" if kind == "bullet" else "ol"
            body.append(f"<{tag}>" + "".join(f"<li>{html_inline(item)}</li>" for item in value) + f"</{tag}>")
        elif kind == "table":
            header, *rows = value
            thead = "".join(f"<th scope=\"col\">{html_inline(cell)}</th>" for cell in header)
            tbody = "".join("<tr>" + "".join(f"<td>{html_inline(cell)}</td>" for cell in row) + "</tr>" for row in rows)
            body.append(f'<div class="table-wrap"><table><thead><tr>{thead}</tr></thead><tbody>{tbody}</tbody></table></div>')

    links = "".join(f'<li><a href="#{anchor}">{html.escape(label)}</a></li>' for anchor, label in toc)
    document = f"""<!doctype html>
<html lang="zh-Hant">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <meta name="description" content="前赤壁賦遊戲式學習企劃書">
  <title>{title}</title>
  <style>
    :root {{ --ink:#1f3029; --muted:#5d6d64; --paper:#f4f2eb; --card:#fffdf7; --green:#425c50; --line:#d7ded6; --soft:#edf1ea; }}
    * {{ box-sizing:border-box; }} html {{ scroll-behavior:smooth; }} body {{ margin:0; background:var(--paper); color:var(--ink); font-family:"Microsoft JhengHei","Noto Sans CJK TC",sans-serif; line-height:1.75; }}
    .layout {{ max-width:1180px; margin:0 auto; padding:36px 22px; display:grid; grid-template-columns:248px minmax(0,1fr); gap:28px; }}
    nav {{ position:sticky; top:18px; align-self:start; max-height:calc(100vh - 36px); overflow:auto; padding:19px; border:1px solid var(--line); border-radius:14px; background:#fff; }}
    nav p {{ margin:0 0 10px; color:var(--muted); font-size:.8rem; font-weight:700; letter-spacing:.08em; }} nav ol {{ margin:0; padding-left:1.25rem; }} nav li {{ margin:.4rem 0; font-size:.9rem; }}
    a {{ color:#315f50; text-decoration-thickness:1px; text-underline-offset:3px; }} nav a {{ text-decoration:none; }} nav a:hover {{ text-decoration:underline; }}
    .hero, article {{ border:1px solid var(--line); border-radius:18px; background:var(--card); }} .hero {{ padding:34px 38px; margin-bottom:20px; }} article {{ padding:28px 38px; }}
    .eyebrow, .meta {{ margin:0; color:var(--muted); }} .eyebrow {{ font-size:.84rem; letter-spacing:.05em; }} h1 {{ margin:.35rem 0; color:#000; font-size:clamp(1.9rem,4.4vw,2.7rem); line-height:1.25; }} .subtitle {{ margin:0; color:var(--green); font-size:1.12rem; font-weight:650; }}
    .summary {{ margin:20px 0 0; padding:14px 17px; border-left:4px solid var(--green); border-radius:6px; background:var(--soft); }}
    h2 {{ margin:2.1rem 0 .8rem; padding-bottom:.32rem; border-bottom:1px solid var(--line); color:#000; font-size:1.4rem; line-height:1.35; scroll-margin-top:18px; }} h2:first-child {{ margin-top:0; }} h3 {{ margin:1.45rem 0 .45rem; color:#000; font-size:1.08rem; }} p {{ margin:.7rem 0; }} ul,ol {{ margin:.65rem 0 1rem; padding-left:1.55rem; }} li {{ margin:.35rem 0; }}
    .note {{ margin:1rem 0; padding:12px 16px; border-left:3px solid #8b9d90; background:#f4f6f2; color:#405047; }} .table-wrap {{ overflow-x:auto; margin:1rem 0; }} table {{ width:100%; border-collapse:collapse; font-size:.94rem; }} th,td {{ padding:10px 12px; border:1px solid var(--line); text-align:left; vertical-align:top; }} th {{ background:var(--green); color:#fff; }} tbody tr:nth-child(even) {{ background:#f5f7f2; }} footer {{ padding:18px; color:var(--muted); text-align:center; font-size:.82rem; }}
    @media(max-width:850px) {{ .layout {{ display:block; padding:16px; }} nav {{ position:static; max-height:none; margin-bottom:16px; }} nav ol {{ columns:2; }} .hero,article {{ padding:24px 20px; }} }} @media(max-width:540px) {{ nav ol {{ columns:1; }} h2 {{ font-size:1.24rem; }} th,td {{ padding:8px; font-size:.9rem; }} }}
    @media print {{ body {{ background:#fff; }} .layout {{ display:block; max-width:none; padding:0; }} nav {{ display:none; }} .hero,article {{ border:0; padding:0 0 18px; }} }}
  </style>
</head>
<body>
  <div class="layout">
    <nav aria-label="企劃書目錄"><p>企劃書目錄</p><ol>{links}</ol></nav>
    <main>
      <header class="hero">
        <p class="eyebrow">遊戲設計文件｜2026 年 10 月 27 日</p>
        <h1>{title}</h1>
        <p class="subtitle">《前赤壁賦》MVP 設計與製作計畫</p>
        <p class="meta">Phaser 4・TypeScript・Vite</p>
        <div class="summary">首版只製作《前赤壁賦》。學生以水墨夜遊場景探索意象與詩句，推論情感並引用原文作證；《桃花源記》待本篇完成及課堂試用後再規劃。</div>
      </header>
      <article>{''.join(body)}</article>
      <footer>月下赤壁 意象尋蹤</footer>
    </main>
  </div>
</body>
</html>
"""
    HTML_OUTPUT.write_text(document, encoding="utf-8")


def set_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = FONT
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    properties = run._element.get_or_add_rPr()
    fonts = properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.insert(0, fonts)
    for attribute in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{attribute}"), FONT)


def style_font(style, size: float, bold: bool, color: str) -> None:
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    properties = style.element.get_or_add_rPr()
    fonts = properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.insert(0, fonts)
    for attribute in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{attribute}"), FONT)


def shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    properties.append(node)


def add_inline(paragraph, value: str, size: float = 10.2, color: str = "26332F") -> None:
    for kind, item in inline_parts(value):
        if kind == "link":
            item = f"{item[0]}（{item[1]}）"
            kind = "text"
        run = paragraph.add_run(item)
        set_font(run, size=size, color=color, bold=(kind == "bold"))
        if kind == "italic":
            run.italic = True


def build_docx(events: list[tuple[str, object]]) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)

    normal = doc.styles["Normal"]
    style_font(normal, 10.2, False, "26332F")
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12
    style_font(doc.styles["Title"], 22, True, "000000")
    style_font(doc.styles["Heading 1"], 15, True, "000000")
    style_font(doc.styles["Heading 2"], 12, True, "000000")
    for name in ("Title", "Heading 1", "Heading 2"):
        doc.styles[name].paragraph_format.keep_with_next = True
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(14)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(5)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(9)
    doc.styles["Heading 2"].paragraph_format.space_after = Pt(4)

    title = doc.add_paragraph(style="Title")
    add_inline(title, "月下赤壁 意象尋蹤 遊戲企劃書", 22, "000000")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(8)
    add_inline(subtitle, "《前赤壁賦》MVP 設計與製作計畫", 11, "405A4B")
    metadata = doc.add_paragraph()
    metadata.paragraph_format.space_after = Pt(13)
    add_inline(metadata, "2026 年 10 月 27 日　｜　Phaser 4・TypeScript・Vite", 8.8, "61706A")

    for kind, value in events:
        if kind == "title":
            continue
        if kind == "heading1":
            paragraph = doc.add_paragraph(style="Heading 1")
            add_inline(paragraph, value, 15, "000000")
        elif kind == "heading2":
            paragraph = doc.add_paragraph(style="Heading 2")
            add_inline(paragraph, value, 12, "000000")
        elif kind == "paragraph":
            paragraph = doc.add_paragraph()
            add_inline(paragraph, value)
        elif kind == "note":
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.16)
            paragraph.paragraph_format.right_indent = Inches(0.12)
            paragraph.paragraph_format.space_after = Pt(8)
            properties = paragraph._p.get_or_add_pPr()
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "EEF2ED")
            properties.append(shading)
            add_inline(paragraph, value, 9.6, "405047")
        elif kind in ("bullet", "number"):
            style = "List Bullet" if kind == "bullet" else "List Number"
            for item in value:
                paragraph = doc.add_paragraph(style=style)
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.keep_together = True
                add_inline(paragraph, item, 10.1)
        elif kind == "table":
            header, *rows = value
            table = doc.add_table(rows=1, cols=len(header))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = "Table Grid"
            for index, cell in enumerate(table.rows[0].cells):
                shade(cell, "425C50")
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                paragraph = cell.paragraphs[0]
                add_inline(paragraph, header[index], 9.2, "FFFFFF")
                for run in paragraph.runs:
                    run.bold = True
            for row_index, row in enumerate(rows):
                cells_row = table.add_row().cells
                for index, cell in enumerate(cells_row):
                    if row_index % 2:
                        shade(cell, "F3F6F1")
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                    add_inline(cell.paragraphs[0], row[index], 9.1)
            spacer = doc.add_paragraph()
            spacer.paragraph_format.space_after = Pt(2)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(footer, "月下赤壁 意象尋蹤", 8, "61706A")
    doc.core_properties.title = "月下赤壁 意象尋蹤 遊戲企劃書"
    doc.core_properties.subject = "前赤壁賦 MVP 設計與製作計畫"
    doc.core_properties.author = ""
    doc.save(DOCX_OUTPUT)


if __name__ == "__main__":
    events = parse_markdown(SOURCE.read_text(encoding="utf-8").splitlines())
    build_html(events)
    build_docx(events)
    print(f"Created {HTML_OUTPUT}")
    print(f"Created {DOCX_OUTPUT}")
